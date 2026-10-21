from sales.models import RecurringSubscription
from datetime import date
from django.contrib.auth.decorators import login_required
from core.decorators import role_required
from django.shortcuts import render, redirect
from django.http import HttpResponse, JsonResponse
from django.core.management import call_command
from inventory.models import Medicine, Batch
from sales.models import SalesInvoice
from finance.services import FinancialService
from inventory.services import AlertService
import io
import openpyxl

@login_required
def dashboard(request):
    alerts = AlertService.get_alerts()
    pnl = FinancialService.get_pnl()
    
    # Pack the actual lists so the template can render them in a modal
    expired_items = []
    for m in alerts['expired']:
        expired_items.append({'name': m['medicine'], 'stock': m['qty'], 'rack': m.get('rack', '')})
        
    near_expiry_items = []
    for m in alerts['near_expiry_30'] + alerts['near_expiry_60'] + alerts['near_expiry_90'] + alerts['near_expiry_180']:
        near_expiry_items.append({'name': m['medicine'], 'stock': m['qty'], 'rack': m.get('rack', '')})
        
    low_stock_items = []
    for m in alerts['low_stock']:
        low_stock_items.append({'name': m['medicine'], 'stock': m['stock'], 'rack': m.get('rack', '')})
    
    # Fetch due subscriptions
    due_subs = RecurringSubscription.objects.filter(is_active=True, next_refill_date__lte=date.today())
    due_subs_list = [{'customer': s.customer.name, 'medicine': s.medicine.name, 'phone': s.customer.phone} for s in due_subs]

    context = {
        'due_subs_list': due_subs_list,
        'due_subs_count': len(due_subs_list),
        'alerts': alerts,
        'pnl': pnl,
        'today_sales': pnl['net_sales_revenue'],
        'today_profit': pnl['net_profit'],
        'low_stock_count': len(alerts['low_stock']),
        'expired_count': len(alerts['expired']),
        'near_expiry_count': len(near_expiry_items),
        'expired_items_json': expired_items,
        'near_expiry_items_json': near_expiry_items,
        'low_stock_items_json': low_stock_items,
        'due_subs_json': due_subs_list,
    }
    return render(request, 'dashboard.html', context)

@login_required
@role_required(['Admin', 'Pharmacist'])
def schedule_register(request):
    from sales.models import SalesItem
    from django.db.models import Q
    
    start_date = parse_date(request.GET.get('start_date'), date.today() - timedelta(days=29))
    end_date = parse_date(request.GET.get('end_date'), date.today())
    
    items = SalesItem.objects.filter(
        batch__medicine__schedule_type__in=['H', 'H1', 'X'],
        invoice__date__date__gte=start_date,
        invoice__date__date__lte=end_date
    ).order_by('-invoice__date')
    
    return render(request, 'schedule_register.html', {
        'items': items,
        'start_date': start_date.strftime('%Y-%m-%d'),
        'end_date': end_date.strftime('%Y-%m-%d')
    })

@login_required
@role_required(['Admin', 'Pharmacist'])
def reports_hub(request):
    return render(request, 'reports_hub.html')

def receipt(request, invoice_id):
    invoice = SalesInvoice.objects.get(id=invoice_id)
    from core.models import PharmacyProfile
    profile = PharmacyProfile.get_profile()
    return render(request, 'receipt.html', {'invoice': invoice, 'profile': profile})

@login_required
@role_required(['Admin', 'Pharmacist'])
def export_stock_excel(request):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=stock.xlsx'
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(['Medicine', 'Batch', 'Expiry', 'Stock (Strips/Tabs)', 'Base Quantity', 'Purchase Price', 'MRP'])
    
    for b in Batch.objects.filter(quantity__gt=0):
        stock_str = ""
        if b.medicine.units_per_pack > 1:
            strips = b.quantity // b.medicine.units_per_pack
            tabs = b.quantity % b.medicine.units_per_pack
            if strips > 0 and tabs > 0:
                stock_str = f"{strips} strips, {tabs} tabs"
            elif strips > 0:
                stock_str = f"{strips} strips"
            else:
                stock_str = f"{tabs} tabs"
        else:
            stock_str = f"{b.quantity} units"
            
        ws.append([b.medicine.name, b.batch_number, b.expiry_date.strftime('%Y-%m-%d'), stock_str, b.quantity, float(b.purchase_price), float(b.mrp)])
        
    wb.save(response)
    return response

@login_required
@role_required(['Admin', 'Pharmacist'])
def export_sales_excel(request):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename=sales.xlsx'
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Invoices"
    ws.append(['Invoice #', 'Date', 'Customer', 'Payment Method', 'Subtotal', 'Discount', 'Net Amount', 'COGS'])
    
    grand_total_net = 0
    for inv in SalesInvoice.objects.all().order_by('-date'):
        cust_name = inv.customer.name if inv.customer else inv.customer_name_manual
        net = float(inv.net_amount)
        grand_total_net += net
        ws.append([
            inv.invoice_number, 
            inv.date.strftime('%Y-%m-%d %H:%M'), 
            cust_name,
            inv.payment_method,
            float(inv.subtotal),
            float(inv.discount),
            net,
            float(inv.total_cogs)
        ])
        
    ws.append(['', '', '', '', '', 'GRAND TOTAL', grand_total_net, ''])
        
    wb.save(response)
    return response

from django.core.cache import cache
from django.contrib.sessions.models import Session
from django.utils import timezone

@login_required
@role_required(['Admin'])
def system_cleanup(request):
    try:
        # Clear Cache
        cache.clear()
        # Clear Expired Sessions
        Session.objects.filter(expire_date__lt=timezone.now()).delete()
        return JsonResponse({'status': 'success', 'message': 'Temporary data and cache cleared successfully.'})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)})

@login_required
@role_required(['Admin'])
def db_backup(request):
    out = io.StringIO()
    call_command('dumpdata', stdout=out)
    response = HttpResponse(out.getvalue(), content_type='application/json')
    response['Content-Disposition'] = 'attachment; filename=backup.json'
    return response

from datetime import timedelta, date, datetime
from reportlab.lib.pagesizes import landscape, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

def parse_date(date_str, default_date):
    if date_str:
        try:
            return datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            return default_date
    return default_date

def generate_pnl_data(start_date=None, end_date=None):
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=29)
    
    data = []
    total_sales = 0
    total_cogs = 0
    total_gross = 0
    total_net = 0
    
    current_date = start_date
    while current_date <= end_date:
        pnl = FinancialService.get_pnl(current_date, current_date)
        if pnl['gross_sales'] > 0 or pnl['total_opex'] > 0 or pnl['expiry_loss'] > 0:
            data.append([
                current_date.strftime('%Y-%m-%d'),
                float(pnl['net_sales_revenue']),
                float(pnl['total_cogs']),
                float(pnl['gross_profit']),
                f"{pnl['gross_margin']:.2f}%",
                float(pnl['total_opex']),
                float(pnl['expiry_loss']),
                float(pnl['net_profit']),
                f"{pnl['net_margin']:.2f}%"
            ])
            total_sales += float(pnl['net_sales_revenue'])
            total_cogs += float(pnl['total_cogs'])
            total_gross += float(pnl['gross_profit'])
            total_net += float(pnl['net_profit'])
        current_date += timedelta(days=1)
        
    return data, total_sales, total_cogs, total_gross, total_net, start_date, end_date

@login_required
@role_required(['Admin'])
def pnl_report(request):
    end_date_str = request.GET.get('end_date')
    start_date_str = request.GET.get('start_date')
    
    end_date = parse_date(end_date_str, date.today())
    start_date = parse_date(start_date_str, end_date - timedelta(days=29))
    
    data, total_sales, total_cogs, total_gross, total_net, s_date, e_date = generate_pnl_data(start_date, end_date)
    
    # overall margins
    total_gross_margin = (total_gross / total_sales * 100) if total_sales > 0 else 0
    total_net_margin = (total_net / total_sales * 100) if total_sales > 0 else 0
    
    context = {
        'start_date': s_date.strftime('%Y-%m-%d'),
        'end_date': e_date.strftime('%Y-%m-%d'),
        'data': data,
        'total_sales': total_sales,
        'total_cogs': total_cogs,
        'total_gross': total_gross,
        'total_net': total_net,
        'total_gross_margin': f"{total_gross_margin:.2f}%",
        'total_net_margin': f"{total_net_margin:.2f}%"
    }
    return render(request, 'pnl_report.html', context)

@login_required
@role_required(['Admin'])
def export_pnl_excel(request):
    start_date = parse_date(request.GET.get('start_date'), date.today() - timedelta(days=29))
    end_date = parse_date(request.GET.get('end_date'), date.today())
    
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename=pnl_report_{start_date}_to_{end_date}.xlsx'
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "P&L Report"
    
    ws.append([
        "Date", "Net Sales Revenue", "COGS", "Gross Profit", 
        "Gross Margin %", "Opex", "Expiry Loss", "Net Profit", "Net Margin %"
    ])
    
    data, total_sales, total_cogs, total_gross, total_net, _, _ = generate_pnl_data(start_date, end_date)
    
    for row in data:
        ws.append(row)
        
    ws.append([
        "GRAND TOTAL",
        round(total_sales, 2),
        round(total_cogs, 2),
        round(total_gross, 2),
        "", "", "",
        round(total_net, 2),
        ""
    ])
    
    wb.save(response)
    return response

@login_required
@role_required(['Admin'])
def export_pnl_pdf(request):
    start_date = parse_date(request.GET.get('start_date'), date.today() - timedelta(days=29))
    end_date = parse_date(request.GET.get('end_date'), date.today())
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="pnl_report_{start_date}_to_{end_date}.pdf"'
    
    doc = SimpleDocTemplate(response, pagesize=landscape(A4))
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph(f"Profit & Loss Report ({start_date} to {end_date})", styles['Title']))
    elements.append(Spacer(1, 12))
    
    table_data = [[
        "Date", "Net Sales", "COGS", "Gross Profit", 
        "Margin %", "Opex", "Expiry Loss", "Net Profit", "Net Margin %"
    ]]
    
    data, total_sales, total_cogs, total_gross, total_net, _, _ = generate_pnl_data(start_date, end_date)
    
    for row in data:
        table_data.append([
            row[0],
            f"Rs {row[1]:.2f}",
            f"Rs {row[2]:.2f}",
            f"Rs {row[3]:.2f}",
            row[4],
            f"Rs {row[5]:.2f}",
            f"Rs {row[6]:.2f}",
            f"Rs {row[7]:.2f}",
            row[8]
        ])
        
    table_data.append([
        "GRAND TOTAL",
        f"Rs {total_sales:.2f}",
        f"Rs {total_cogs:.2f}",
        f"Rs {total_gross:.2f}",
        "", "", "",
        f"Rs {total_net:.2f}",
        ""
    ])
    
    t = Table(table_data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('BACKGROUND', (0,-1), (-1,-1), colors.lightgrey),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 1, colors.black)
    ]))
    
    elements.append(t)
    doc.build(elements)
    return response

@login_required
@role_required(['Admin', 'Pharmacist'])
def export_stock_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="stock_report.pdf"'
    
    doc = SimpleDocTemplate(response, pagesize=landscape(A4))
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("Current Inventory / Stock Report", styles['Title']))
    elements.append(Spacer(1, 12))
    
    table_data = [["Medicine", "Batch No.", "Expiry", "Stock (Strips/Tabs)", "Purchase Price", "MRP"]]
    
    for b in Batch.objects.filter(quantity__gt=0).order_by('medicine__name', 'expiry_date'):
        stock_str = ""
        if b.medicine.units_per_pack > 1:
            strips = b.quantity // b.medicine.units_per_pack
            tabs = b.quantity % b.medicine.units_per_pack
            if strips > 0 and tabs > 0:
                stock_str = f"{strips} strips, {tabs} tabs"
            elif strips > 0:
                stock_str = f"{strips} strips"
            else:
                stock_str = f"{tabs} tabs"
        else:
            stock_str = f"{b.quantity} units"
            
        table_data.append([
            b.medicine.name,
            b.batch_number,
            b.expiry_date.strftime('%Y-%m-%d'),
            stock_str,
            f"Rs {b.purchase_price:.2f}",
            f"Rs {b.mrp:.2f}"
        ])
        
    t = Table(table_data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('ALIGN', (3,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('GRID', (0,0), (-1,-1), 1, colors.black)
    ]))
    
    elements.append(t)
    doc.build(elements)
    return response

@login_required
@role_required(['Admin', 'Pharmacist'])
def export_sales_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="sales_report.pdf"'
    
    doc = SimpleDocTemplate(response, pagesize=landscape(A4))
    elements = []
    styles = getSampleStyleSheet()
    
    elements.append(Paragraph("Sales Report", styles['Title']))
    elements.append(Spacer(1, 12))
    
    table_data = [["Invoice #", "Date", "Customer", "Subtotal", "Discount", "Net Amount", "COGS"]]
    
    grand_total_net = 0
    for inv in SalesInvoice.objects.all().order_by('-date'):
        cust_name = inv.customer.name if inv.customer else inv.customer_name_manual
        net = float(inv.net_amount)
        grand_total_net += net
        table_data.append([
            inv.invoice_number,
            inv.date.strftime('%Y-%m-%d %H:%M'),
            cust_name,
            f"Rs {inv.subtotal:.2f}",
            f"Rs {inv.discount:.2f}",
            f"Rs {net:.2f}",
            f"Rs {inv.total_cogs:.2f}"
        ])
        
    table_data.append([
        "GRAND TOTAL", "", "", "", "", f"Rs {grand_total_net:.2f}", ""
    ])
    
    t = Table(table_data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('ALIGN', (3,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('BACKGROUND', (0,-1), (-1,-1), colors.lightgrey),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('GRID', (0,0), (-1,-1), 1, colors.black)
    ]))
    
    elements.append(t)
    doc.build(elements)
    return response

